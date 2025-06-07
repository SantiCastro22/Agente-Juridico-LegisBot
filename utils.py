import re
from docx import Document
import os

def extraer_datos_cliente(docx_path):
    """
    Extrae datos clave-valor de un DOCX tipo formulario (una línea por dato, formato: Campo: Valor)
    """
    doc = Document(docx_path)
    datos = {}
    for para in doc.paragraphs:
        match = re.match(r"([\w\sáéíóúñÁÉÍÓÚÑ]+):\s*(.+)", para.text)
        if match:
            campo = match.group(1).strip().lower().replace(" ", "_")
            valor = match.group(2).strip()
            datos[campo] = valor
    return datos

def reemplazar_placeholders(texto, datos):
    """
    Reemplaza {campo} en texto por el valor correspondiente en datos.
    """
    def repl(match):
        key = match.group(1).lower().replace(" ", "_")
        return datos.get(key, match.group(0))
    return re.sub(r"{([\w\sáéíóúñÁÉÍÓÚÑ]+)}", repl, texto)